"""Objetos de Diosesmon: cómo se consigue cada uno y si su crafteo está activo en el servidor.

  python -I tools/build_objetos_data.py --crear   Crea tools/source/objetos-diosesmon.xlsx o, si ya existe, le añade
                                                  los objetos nuevos conservando lo que hayas escrito.
  python -I tools/build_objetos_data.py           Lee el Excel y genera data/objetos-data.js para la página Objetos.

Fuentes:
  - Cobblemon 1.7.3 (tools/build-cache/cobblemon-1.7.3-data.tar.gz y el .jar de la carpeta mods): recetas, etiquetas,
    tablas de botín, nombres y descripciones en español y texturas.
  - Mods Mega Showdown y Navas ZA Megas (sus .jar en la carpeta mods): objetos, recetas, nombres y texturas.
  - Datos de la wiki: kits, tesoros, drops de Pokémon (respetando drop-overrides.js), cazas, gimnasios, DexRewards
    y la tienda de la Zona Safari.
En el Excel (hoja «Objetos») el staff marca «Crafteo en Diosesmon» (Activo / Desactivado / Sin receta) y puede
escribir «Se consigue en» y «Nota». La carpeta de mods se puede pasar con --mods <carpeta>.
"""
import json
import re
import site
import sys
import tarfile
import unicodedata
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

# Con «python -I» no se carga la carpeta de paquetes del usuario, que es donde suele estar openpyxl.
if site.getusersitepackages() not in sys.path:
    sys.path.append(site.getusersitepackages())
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from diosesmon_config import is_available  # noqa: E402

CACHE = ROOT / "tools" / "build-cache"
EXCEL = ROOT / "tools" / "source" / "objetos-diosesmon.xlsx"
# Recetas reales del servidor, sacadas en el juego con el mod diosesmon-recetas (/wikirecetas).
# Se actualizan con:  python -I tools/build_objetos_data.py --recetas "<instancia>/diosesmon-recetas.json"
SERVER_RECIPES = ROOT / "tools" / "source" / "recetas-servidor.json"
OUTPUT = ROOT / "data" / "objetos-data.js"
IMG = ROOT / "assets" / "img" / "items"
MODS = Path.home() / "curseforge/minecraft/Instances/Diosesmon Oficial - Lite/mods"
if "--mods" in sys.argv:
    MODS = Path(sys.argv[sys.argv.index("--mods") + 1])
MINECRAFT_TEXTURES = "https://raw.githubusercontent.com/InventivetalentDev/minecraft-assets/1.21.1/assets/minecraft/textures"

ORIGINS = {"cobblemon": "Cobblemon", "mega_showdown": "Mega Showdown", "zamega": "Navas ZA Megas", "diosesmon": "Diosesmon"}
STATES = {"activo": "on", "desactivado": "off", "sin receta": "none"}
STATE_LABEL = {"on": "Activo", "off": "Desactivado", "none": "Sin receta"}
HEADERS = ["ID", "Nombre", "Categoría", "Origen", "Tiene receta en el mod", "Crafteo en Diosesmon", "Se consigue en", "Nota"]

CATEGORIES = [
    ("balls", "Poké Balls", "⚪"),
    ("medicina", "Medicinas", "💊"),
    ("evolucion", "Objetos de evolución", "✨"),
    ("combate", "Equipables y combate", "⚔️"),
    ("entrenamiento", "Vitaminas, EVs y caramelos", "💪"),
    ("crianza", "Crianza, naturaleza y habilidad", "🥚"),
    ("megas", "Megapiedras y Megapulsera", "💎"),
    ("zdt", "Movimientos Z, Dinamax y Teracristal", "🌟"),
    ("formas", "Objetos de formas y legendarios", "🔮"),
    ("cocina", "Cocina, ingredientes y cultivo", "🍲"),
    ("fosiles", "Fósiles y arqueología", "🦴"),
    ("bloques", "Bloques y máquinas", "🧱"),
    ("otros", "Herramientas y otros", "🎒"),
]

# Reglas ya conocidas de Diosesmon: solo se usan al crear el Excel (después manda lo que diga el Excel).
KNOWN_DISABLED = {
    "cobblemon:poke_snack": "Los Poké Snacks están desactivados en Diosesmon.",
}
IV_CANDY_NOTE = ("Los caramelos de IV están desactivados en Diosesmon. Para mejorar IVs usa la crianza "
                 "o el NPC Tienda de IVs.")

# Objetos propios del servidor (no son de Cobblemon): la tienda de la Zona Safari.
SAFARI_EXTRA = [
    ("diosesmon:jet_ball", "Jet Ball", "balls", "Ratio de captura ×2 y vuela más lejos. Solo vale dentro de la Zona Safari.", "assets/img/safari-tienda/jet-ball.png"),
    ("diosesmon:bola_fango", "Bola de fango", "otros", "Lánzala a un Pokémon de la Zona Safari: puede dejarlo aturdido.", "assets/img/safari-tienda/bola-fango.png"),
    ("diosesmon:bola_pegajosa", "Bola pegajosa", "otros", "Lánzala a un Pokémon de la Zona Safari: siempre lo deja aturdido.", "assets/img/safari-tienda/bola-pegajosa.png"),
    ("diosesmon:bomba_humo", "Bomba de humo", "otros", "Úsala en la Zona Safari para que los Pokémon tarden en detectarte.", "assets/img/safari-tienda/bomba-humo.png"),
]
# Objetos internos o de administración que no se muestran.
HIDDEN = {"cobblemon:npc_editor", "mega_showdown:debug_stick"}

# Nombres legibles de las etiquetas que salen en las recetas: (nombre, objeto que se usa de icono).
TAG_LABELS = {
    "c:ingots/iron": ("Lingote de hierro", "minecraft:iron_ingot"),
    "c:ingots/copper": ("Lingote de cobre", "minecraft:copper_ingot"),
    "c:ingots/gold": ("Lingote de oro", "minecraft:gold_ingot"),
    "c:ingots/netherite": ("Lingote de netherita", "minecraft:netherite_ingot"),
    "c:nuggets/iron": ("Pepita de hierro", "minecraft:iron_nugget"),
    "c:nuggets/gold": ("Pepita de oro", "minecraft:gold_nugget"),
    "c:gems/amethyst": ("Fragmento de amatista", "minecraft:amethyst_shard"),
    "c:gems/diamond": ("Diamante", "minecraft:diamond"),
    "c:gems/quartz": ("Cuarzo del Nether", "minecraft:quartz"),
    "c:gems/lapis": ("Lapislázuli", "minecraft:lapis_lazuli"),
    "c:gems/prismarine": ("Cristal de prismarina", "minecraft:prismarine_crystals"),
    "c:chains": ("Cadena", "minecraft:chain"),
    "c:strings": ("Hilo", "minecraft:string"),
    "c:dusts/redstone": ("Polvo de redstone", "minecraft:redstone"),
    "c:chests/wooden": ("Cofre de madera", "minecraft:chest"),
    "c:concretes": ("Cualquier hormigón", "minecraft:white_concrete"),
    "c:leathers": ("Cuero", "minecraft:leather"),
    "c:slime_balls": ("Bola de slime", "minecraft:slime_ball"),
    "c:drinks/milk": ("Leche (cubo o Leche Mu-Mu)", "minecraft:milk_bucket"),
    "c:fertilizers": ("Fertilizante (harina de huesos)", "minecraft:bone_meal"),
    "c:tools/shield": ("Escudo", "minecraft:shield"),
    "c:rods/wooden": ("Palo", "minecraft:stick"),
    "c:rods/blaze": ("Vara de blaze", "minecraft:blaze_rod"),
    "c:storage_blocks/iron": ("Bloque de hierro", "minecraft:iron_block"),
    "c:raw_materials/gold": ("Oro en bruto", "minecraft:raw_gold"),
    "c:seeds": ("Cualquier semilla", "minecraft:wheat_seeds"),
    "c:crops/wheat": ("Trigo", "minecraft:wheat"),
    "c:bricks/normal": ("Ladrillo", "minecraft:brick"),
    "c:buckets/empty": ("Cubo", "minecraft:bucket"),
    "c:foods/raw_meat": ("Cualquier carne cruda", "minecraft:beef"),
    "c:mushrooms": ("Cualquier champiñón", "minecraft:red_mushroom"),
    "c:foods/bread": ("Pan", "minecraft:bread"),
    "minecraft:enchantable/fishing": ("Caña de pescar o Poké Caña", "minecraft:fishing_rod"),
    "minecraft:wool": ("Cualquier lana", "minecraft:white_wool"),
    "minecraft:chicken_food": ("Semillas (comida de gallina)", "minecraft:wheat_seeds"),
    "minecraft:trim_materials": ("Material de adorno", "minecraft:iron_ingot"),
    "minecraft:trimmable_armor": ("Pieza de armadura", "minecraft:iron_chestplate"),
    "minecraft:buttons": ("Cualquier botón", "minecraft:oak_button"),
    "minecraft:planks": ("Cualquier tablón", "minecraft:oak_planks"),
    "minecraft:wooden_slabs": ("Cualquier losa de madera", "minecraft:oak_slab"),
    "cobblemon:tier_1_poke_ball_materials": ("Material de Poké Ball (hierro)", "minecraft:iron_ingot"),
    "cobblemon:tier_2_poke_ball_materials": ("Material de Super Ball (oro)", "minecraft:gold_ingot"),
    "cobblemon:tier_3_poke_ball_materials": ("Material de Ultra Ball (diamante)", "minecraft:diamond"),
    "cobblemon:tier_4_poke_ball_materials": ("Material de Master Ball (netherita)", "minecraft:netherite_ingot"),
    "cobblemon:plaques": ("Cualquier placa", None),
    "cobblemon:pokedex_screen": ("Pantalla de Pokédex (cristal tintado)", None),
    "cobblemon:berries": ("Cualquier baya", None),
    "cobblemon:remedy_berries": ("Baya curativa", None),
    "cobblemon:super_potion_ingredients": ("Ingrediente de Superpoción", None),
    "cobblemon:apricorns": ("Cualquier bonguri", None),
    "cobblemon:apricorn_logs": ("Tronco de bonguri", None),
    "cobblemon:saccharine_logs": ("Tronco de sacarino", None),
    "cobblemon:full_heal_bottles": ("Botella de Cura Total", None),
    "cobblemon:full_heal_ingredients": ("Ingrediente de Cura Total", None),
    "cobblemon:apples": ("Cualquier manzana", None),
    "cobblemon:sandwich_veggies": ("Verdura para sándwich", None),
}
DYE_COLORS = {"black": "negro", "blue": "azul", "brown": "marrón", "cyan": "cian", "gray": "gris", "green": "verde",
              "light_blue": "azul claro", "light_gray": "gris claro", "lime": "verde lima", "magenta": "magenta", "orange": "naranja",
              "pink": "rosa", "purple": "morado", "red": "rojo", "white": "blanco", "yellow": "amarillo"}
# Iconos de respaldo (capturas del juego que ya usa la wiki) para bloques con modelo 3D.
ICON_FALLBACK = {
    "cobblemon:healing_machine": "assets/img/tesoros/gacha-diario/maquina-de-curacion.png",
    "cobblemon:pc": "assets/img/tesoros/gacha-diario/pc.png",
}
# Nombres que el mod deja en inglés o mal traducidos.
NAME_FIXES = {"mega_showdown:gracidea_flower": "Gracídea"}
# Texturas que no se llaman como el objeto.
TEXTURE_ALIASES = {
    "cobblemon:x_defence": "x_defense", "cobblemon:x_special_attack": "x_sp_atk", "cobblemon:x_special_defence": "x_sp_def",
    "cobblemon:ponigiri": "ponigiri_base", "cobblemon:display_case": "display_case_side", "cobblemon:hearty_grain_bale": "hearty_grains",
}
# Iconos de Minecraft que no son item/<nombre>.png.
MC_ICONS = {
    "minecraft:clock": "item/clock_00", "minecraft:composter": "block/composter_side",
    "minecraft:flowering_azalea": "block/flowering_azalea_side", "minecraft:lilac": "block/lilac_top",
    "minecraft:peony": "block/peony_top", "minecraft:rose_bush": "block/rose_bush_top", "minecraft:sunflower": "block/sunflower_front",
    "minecraft:magma_block": "block/magma", "minecraft:melon": "block/melon_side", "minecraft:oak_button": "block/oak_planks",
    "minecraft:oak_slab": "block/oak_planks", "minecraft:piston": "block/piston_side", "minecraft:smooth_stone_slab": "block/smooth_stone_slab_side",
    "minecraft:target": "block/target_side",
}
# Recompensas que se escriben distinto en otras páginas.
NAME_ALIASES = {
    "caramelo xl": "cobblemon:exp_candy_xl", "caramelos xl": "cobblemon:exp_candy_xl",
    "repartir exp": "cobblemon:exp_share", "gracidea": "mega_showdown:gracidea_flower",
    "maquina de curacion": "cobblemon:healing_machine", "pc": "cobblemon:pc",
    "prisma rojo": "mega_showdown:red_orb", "prisma azul": "mega_showdown:blue_orb",
    "punta adn": "mega_showdown:dna_splicer", "gran griseosfera": "mega_showdown:griseous_core",
    "gran diamansfera": "mega_showdown:adamant_crystal", "gran lustresfera": "mega_showdown:lustrous_globe",
    "vasija castigo": "mega_showdown:prison_bottle", "megapulsera": "mega_showdown:mega_bracelet",
}


def fold(text):
    text = unicodedata.normalize("NFD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", text)).strip()


def clean(text):
    return re.sub(r"§.", "", str(text or "")).strip()


def js_value(path, name):
    source = (ROOT / path).read_text(encoding="utf-8")
    start = source.index(f"window.{name} = ") + len(f"window.{name} = ")
    try:
        return json.JSONDecoder().raw_decode(source[start:])[0]
    except json.JSONDecodeError:
        return None


def js_object_loose(path, name):
    """Lee objetos JS escritos a mano (claves sin comillas, comentarios, comas finales)."""
    source = (ROOT / path).read_text(encoding="utf-8")
    start = source.index(f"window.{name} = ") + len(f"window.{name} = ")
    body = source[start:source.rindex("];") + 1] if source[start] == "[" else source[start:source.rindex("}") + 1]
    body = re.sub(r"^\s*//.*$", "", body, flags=re.M)
    body = re.sub(r'(?<![:"\w])//[^\n"]*$', "", body, flags=re.M)
    body = re.sub(r"([{,]\s*)([A-Za-z_]\w*)\s*:", r'\1"\2":', body)
    body = re.sub(r",(\s*[}\]])", r"\1", body)
    return json.loads(body)


def find_jar(prefix):
    jars = sorted(MODS.glob(prefix + "*.jar"))
    if not jars:
        raise SystemExit(f"No encuentro {prefix}*.jar en {MODS} (usa --mods <carpeta>)")
    return zipfile.ZipFile(jars[-1])


# ---------------------------------------------------------------- Datos de los mods

def load_mod_data():
    lang = json.loads((CACHE / "cobblemon-1.7.3-es_es.json").read_text(encoding="utf-8-sig"))
    mc_lang = json.loads((CACHE / "minecraft-1.21.1-es_es.json").read_text(encoding="utf-8-sig"))
    jars = {"cobblemon": find_jar("Cobblemon-fabric"), "mega_showdown": find_jar("mega_showdown-fabric"), "zamega": find_jar("zamega-fabric")}
    en = {}
    for jar in jars.values():
        for name in jar.namelist():
            if name.endswith("lang/es_es.json"):
                lang.update(json.loads(jar.read(name).decode("utf-8-sig")))
            elif name.endswith("lang/en_us.json"):
                en.update(json.loads(jar.read(name).decode("utf-8-sig")))

    recipes, tags, loot = [], {}, {}
    with tarfile.open(CACHE / "cobblemon-1.7.3-data.tar.gz") as tar:
        for member in tar.getmembers():
            if not member.isfile() or not member.name.endswith(".json") or "/data/" not in member.name:
                continue
            path = member.name.split("/data/", 1)[1]
            if path.startswith("cobblemon/recipe/"):
                recipes.append(("cobblemon", path, json.loads(tar.extractfile(member).read())))
            elif path.startswith("cobblemon/tags/item/"):
                tags["cobblemon:" + path[len("cobblemon/tags/item/"):-5]] = json.loads(tar.extractfile(member).read()).get("values", [])
            elif path.startswith("cobblemon/loot_table/"):
                loot["cobblemon:" + path[len("cobblemon/loot_table/"):-5]] = json.loads(tar.extractfile(member).read())
    for key in ("mega_showdown", "zamega"):
        jar = jars[key]
        for name in jar.namelist():
            match = re.match(r"data/([^/]+)/(recipe|tags/item|loot_table)/(.+)\.json$", name)
            if not match:
                continue
            ns, kind, rest = match.groups()
            data = json.loads(jar.read(name))
            if kind == "recipe":
                recipes.append((key, name, data))
            elif kind == "tags/item":
                tags.setdefault(f"{ns}:{rest}", []).extend(data.get("values", []))
            else:
                loot[f"{ns}:{rest}"] = data
    return lang, mc_lang, en, jars, recipes, tags, loot


def resolve_tag(tags, tag, seen=None):
    seen = seen or set()
    out = []
    for value in tags.get(tag, []):
        value = value if isinstance(value, str) else value.get("id", "")
        if value.startswith("#"):
            if value[1:] not in seen:
                seen.add(value[1:])
                out += resolve_tag(tags, value[1:], seen)
        else:
            out.append(value)
    return out


def loot_items(loot, key, seen=None):
    """Objetos que puede dar una tabla de botín (siguiendo las referencias a otras tablas)."""
    seen = seen or set()
    if key in seen or key not in loot:
        return set()
    seen.add(key)
    found = set()

    def walk(node):
        if isinstance(node, dict):
            kind = node.get("type", "")
            if kind == "minecraft:item" and isinstance(node.get("name"), str):
                found.add(node["name"])
            elif kind == "minecraft:loot_table":
                ref = node.get("value") or node.get("name")
                if isinstance(ref, str):
                    found.update(loot_items(loot, ref, seen))
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)
    walk(loot[key])
    return found


# ---------------------------------------------------------------- Catálogo

def category_of(item_id, tags_of, is_block):
    ns, _, path = item_id.partition(":")
    has = lambda tag: tag in tags_of.get(item_id, set())
    if ns in ("mega_showdown", "zamega"):
        if ns == "zamega" or has("mega_showdown:mega_stone") or has("mega_showdown:mega_bracelet") or path in {
                "mega_stone", "keystone", "mega_ring", "brendan_mega_cuff", "may_bracelet", "lysandre_ring", "maxie_glasses",
                "archie_anchor", "lisia_mega_tiara", "korrina_glove", "red_orb", "blue_orb", "redorb", "blueorb"} or path.startswith("mega_bracelet"):
            return "megas"
        if path.endswith("_ore") or path.endswith("_block") or "meteorid" in path or "mega_evo" in path or path in {"pedestal", "mega_stone_crystal"}:
            return "bloques"
        if (has("mega_showdown:z_crystal") or has("mega_showdown:z_ring") or has("mega_showdown:tera_shard") or path.endswith("_z")
                or "z_ring" in path or "z_power_ring" in path or path.startswith("tera_") or path.startswith("max_") or "dynamax" in path
                or "wishing_star" in path or path in {"omni_ring", "sweet_max_soup", "power_spot", "star_core", "sparkling_stone_light", "sparkling_stone_dark", "dormant_crystal", "likos_pendant"}):
            return "zdt"
        return "formas"
    if has("cobblemon:berries"):
        return "bayas"
    if has("cobblemon:poke_balls"):
        return "balls"
    if (has("cobblemon:potions") or has("cobblemon:restores") or has("cobblemon:revives") or has("cobblemon:remedies") or has("cobblemon:ethers")
            or path in {"heal_powder", "energy_root", "revival_herb", "berry_juice", "medicinal_brew", "full_heal", "medicinal_leek"}):
        return "medicina"
    if path in {"destiny_knot", "everstone", "ability_capsule", "ability_patch"} or has("cobblemon:mints") or has("cobblemon:mint_leaves") or has("cobblemon:mint_seeds"):
        return "crianza"
    if (has("cobblemon:vitamins") or has("cobblemon:feathers") or has("cobblemon:mochis") or has("cobblemon:candies") or path == "rare_candy"
            or path.startswith("power_") and path in {"power_anklet", "power_band", "power_belt", "power_bracer", "power_lens", "power_weight"}
            or path in {"exp_share", "lucky_egg", "macho_brace"}):
        return "entrenamiento"
    if has("cobblemon:evolution_items") or has("cobblemon:evolution_stones"):
        return "evolucion"
    if has("cobblemon:fossils") or "sherd" in path or path in {"relic_coin", "relic_coin_pouch", "relic_coin_sack"} or "tumblestone" in path and not is_block:
        return "fosiles"
    if has("cobblemon:held/is_held_item") or has("cobblemon:battle_items") or has("cobblemon:type_gems") or has("cobblemon:held/terrain_seeds"):
        return "combate"
    if (has("cobblemon:poke_food") or has("cobblemon:aprijuices") or has("cobblemon:apricorns") or has("cobblemon:sweets") or has("cobblemon:seeds")
            or has("cobblemon:apricorn_sprouts") or "mulch" in path or "seed" in path or path in {
                "vivichoke", "hearty_grains", "galarica_nuts", "bugwort", "moomoo_milk", "tasty_tail", "poke_bait", "poke_puff", "roasted_leek",
                "braised_vivichoke", "vivichoke_dip", "leek_and_potato_stew", "ponigiri", "fried_rice", "potato_mochi", "jelly_doughnut", "cream_puff",
                "poke_cake", "lure_cake", "poke_snack"} or path.endswith("_apricorn")):
        return "cocina"
    if is_block or has("cobblemon:machines") or "boat" in path or "plaque" in path or "campfire_pot" in path:
        return "bloques"
    return "otros"


def description(lang, item_id, comp_effects, stone_info):
    ns, _, path = item_id.partition(":")
    if item_id in stone_info:
        return stone_info[item_id]
    if ns == "cobblemon" and path in comp_effects:
        return comp_effects[path]
    keys = sorted(k for k in lang if k.startswith(f"item.{ns}.{path}.") and ("tooltip" in k or k.endswith(".desc") or k.endswith(".description")))
    text = " ".join(clean(lang[k]) for k in keys if clean(lang[k]))
    if not text:
        text = clean(lang.get(f"tooltip.{ns}.{path}.tooltip") or lang.get(f"block.{ns}.{path}.tooltip") or lang.get(f"block.{ns}.{path}.desc") or "")
    if not text:
        text = next((value for pattern, value in DEFAULT_DESCRIPTIONS if re.search(pattern, path)), "")
    return re.sub(r"\s+", " ", text)[:400]


# Descripciones por defecto cuando el mod no trae ninguna (patrón del id → texto).
DEFAULT_DESCRIPTIONS = [
    (r"_tera_shard$", "Teralito: con 50 del mismo tipo se cambia el teratipo de un Pokémon."),
    (r"^tera_pouch", "Bolsa para guardar Teralitos."),
    (r"(ium|ium_z|_z)$", "Cristal Z: si lo lleva un Pokémon y su entrenador tiene un Aro Z, puede usar un movimiento Z."),
    (r"z_power_ring$|z_ring", "Aro Z: permite a tus Pokémon usar movimientos Z si llevan un cristal Z."),
    (r"_plate$", "Tabla: si la lleva Arceus cambia su tipo (y el de Sentencia)."),
    (r"_memory$", "Disco de memoria: si lo lleva Silvally cambia su tipo."),
    (r"_drive$", "ROM: si la lleva Genesect cambia el tipo de Tecno Rayo."),
    (r"^rotom_", "Electrodoméstico: Rotom cambia de forma al entrar en él."),
    (r"_nectar$", "Néctar: Oricorio cambia de estilo al tomarlo."),
    (r"_mask$", "Máscara de Ogerpon: cambia su forma y su tipo."),
    (r"^max_|^sweet_max_soup$|^dynamax_candy$", "Objeto del Dinamax (Mega Showdown)."),
    (r"_fossil$|^fossilized_", "Fósil: revívelo en la Máquina Restauradora."),
    (r"_sherd$", "Fragmento de cerámica: decora vasijas. Se encuentra excavando bloques sospechosos."),
    (r"mulch", "Abono para los árboles de bayas: mira la sección Bayas."),
    (r"^aprijuice_", "Zumo de bonguri: comida que mejora las estadísticas de montura de un Pokémon."),
    (r"_apricorn_seed$", "Semilla de bonguri: plántala para que crezca un árbol de bonguris."),
    (r"_apricorn$", "Bonguri: se usa para fabricar Poké Balls y zumos."),
    (r"_rod$", "Poké Caña: caña de pescar Pokémon que usa esta Poké Ball como corcho."),
    (r"^pokedex_", "Pokédex: registra los Pokémon que ves y capturas."),
    (r"_smithing_template$", "Plantilla de herrería."),
    (r"^(sky_|black_)?tumblestone$", "Piedra rodada: material para fabricar Poké Balls antiguas."),
    (r"relic_coin", "Moneda Reliquia: Gimmighoul y Gholdengo las coleccionan."),
]


def texture_index(jar, namespace):
    paths = {}
    for name in jar.namelist():
        prefix = f"assets/{namespace}/textures/"
        if not name.startswith(prefix) or not name.endswith(".png"):
            continue
        sub = name[len(prefix):]
        if not (sub.startswith("item/") or sub.startswith("block/")):
            continue
        stem = Path(name).stem
        score = (0 if sub.startswith("item/") else 2) + (1 if "/models/" in sub or "/model/" in sub else 0)
        if stem not in paths or score < paths[stem][0]:
            paths[stem] = (score, name)
    return {stem: value[1] for stem, value in paths.items()}


def ensure_icon(item_id, jars, indexes):
    """Copia o descarga el icono si falta. Devuelve la ruta relativa o None."""
    ns, _, path = item_id.partition(":")
    target = IMG / ns / f"{path}.png"
    if target.exists():
        return target.relative_to(ROOT).as_posix()
    data = None
    base = re.sub(r"_(slab|stairs|wall|fence_gate|fence|button|pressure_plate|wood)$", "", path)
    candidates = [TEXTURE_ALIASES.get(item_id, path), path + "_side", path + "_front", path + "_top"]
    if base != path:
        candidates += [base + "_planks", base + "s", base, base + "_log"]
    if ns in jars:
        index = indexes.setdefault(ns, texture_index(jars[ns], ns))
        for candidate in candidates:
            if candidate in index:
                data = jars[ns].read(index[candidate])
                width, height = int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
                if width > height or width >= 64:
                    data = None  # Textura de modelo 3D (máquinas, cofres…): no sirve de icono.
                    continue
                break
    elif ns == "minecraft":
        urls = [f"{MINECRAFT_TEXTURES}/{MC_ICONS[item_id]}.png"] if item_id in MC_ICONS else []
        urls += [f"{MINECRAFT_TEXTURES}/{folder}/{path}.png" for folder in ("item", "block")]
        for url in urls:
            try:
                data = urllib.request.urlopen(url, timeout=30).read()
                break
            except (urllib.error.HTTPError, urllib.error.URLError):
                continue
    if not data:
        return None
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return target.relative_to(ROOT).as_posix()


def build_catalog():
    lang, mc_lang, en, jars, raw_recipes, tags, loot = load_mod_data()
    tags_of = {}
    for tag in tags:
        for value in resolve_tag(tags, tag):
            tags_of.setdefault(value, set()).add(tag)

    # Todos los objetos: claves de idioma + resultados de recetas, de los tres mods.
    names, blocks = {}, set()
    for source in (en, lang):
        for key, value in source.items():
            match = re.fullmatch(r"(item|block)\.(cobblemon|mega_showdown|zamega)\.([a-z0-9_]+)", key)
            if not match:
                continue
            kind, ns, path = match.groups()
            item_id = f"{ns}:{path}"
            if path.startswith("potted_") or path.endswith("_wall_sign") or path.endswith("_wall_hanging_sign"):
                continue
            if kind == "item" or item_id not in names:
                names[item_id] = NAME_FIXES.get(item_id) or clean(value)
            if kind == "block":
                blocks.add(item_id)
    items_with_item_key = {k.split(".", 1)[1].replace(".", ":", 1) for k in list(lang) + list(en) if k.startswith("item.")}
    blocks -= items_with_item_key
    # Bloques que no son objetos (cultivos, variantes colgadas…): solo si tienen textura de objeto o receta.
    def es_name(item_id):
        ns, _, path = item_id.partition(":")
        if item_id in NAME_FIXES:
            return NAME_FIXES[item_id]
        ore = re.fullmatch(r"(deepslate_|nether_|dripstone_|terracotta_|mega_meteorid_)?(\w+?)(_stone)?_ore", path)
        if ore and ns in ("cobblemon", "mega_showdown") and (ore.group(3) or ore.group(1) == "mega_meteorid_"):
            stone = es_name(f"cobblemon:{ore.group(2)}_stone")
            where = {"deepslate_": " de pizarra profunda", "nether_": " del Nether", "dripstone_": " de espeleotema",
                     "terracotta_": " de terracota", "mega_meteorid_": " de meteorito"}.get(ore.group(1) or "", "")
            return f"Mena de {stone}{where}"
        if ns == "minecraft":
            return clean(mc_lang.get(f"item.minecraft.{path}") or mc_lang.get(f"block.minecraft.{path}") or path.replace("_", " ").capitalize())
        return clean(lang.get(f"item.{ns}.{path}") or lang.get(f"block.{ns}.{path}") or names.get(item_id) or path.replace("_", " ").capitalize())

    # Recetas por objeto resultado.
    def ref(ingredient):
        if isinstance(ingredient, list):
            options = [ref(option) for option in ingredient]
            return "|".join(option for option in options if option)
        if not isinstance(ingredient, dict):
            return str(ingredient) if ingredient else None
        if "tag" in ingredient:
            return "#" + ingredient["tag"]
        return ingredient.get("item") or ingredient.get("id")

    kinds = {"minecraft:crafting_shaped": "shaped", "minecraft:crafting_shapeless": "shapeless", "cobblemon:cooking_pot": "pot",
             "cobblemon:cooking_pot_shapeless": "potless", "minecraft:smelting": "smelt", "minecraft:blasting": "blast",
             "minecraft:smoking": "smoke", "minecraft:campfire_cooking": "campfire", "cobblemon:brewing_stand": "brew",
             "minecraft:stonecutting": "stone", "minecraft:smithing_transform": "smith"}
    recipes = {}
    for mod, path, data in raw_recipes:
        kind = kinds.get(data.get("type"))
        result = data.get("result")
        result_id = result.get("id") if isinstance(result, dict) else result
        if not kind or not isinstance(result_id, str):
            continue
        count = result.get("count", 1) if isinstance(result, dict) else data.get("count", 1)
        recipe = {"k": kind, "m": mod}
        if count and count > 1:
            recipe["n"] = count
        if kind in ("shaped", "pot"):
            pattern = data.get("pattern", [])
            width = max(len(row) for row in pattern) if pattern else 0
            grid = []
            for row in pattern:
                grid.append([ref(data["key"][ch]) if ch != " " and ch in data.get("key", {}) else None for ch in row.ljust(width)])
            recipe["g"] = grid
        elif kind in ("shapeless", "potless"):
            counted = {}
            for ingredient in data.get("ingredients", []):
                key = ref(ingredient)
                counted[key] = counted.get(key, 0) + 1
            recipe["i"] = [[key, amount] for key, amount in counted.items()]
        elif kind == "brew":
            recipe["i"] = [[ref(data.get("bottle")), 1], [ref(data.get("input")), 1]]
        elif kind == "smith":
            recipe["i"] = [[ref(data.get(part)), 1] for part in ("template", "base", "addition") if data.get(part)]
        else:
            recipe["i"] = [[ref(data.get("ingredient")), 1]]
        if "seasoningTag" in data and data["seasoningTag"] not in ("cobblemon:empty", None):
            recipe["s"] = True
        recipes.setdefault(result_id, [])
        if recipe not in recipes[result_id]:
            recipes[result_id].append(recipe)
    for result_id in recipes:
        if result_id.split(":")[0] in ORIGINS and result_id not in names:
            names[result_id] = es_name(result_id)

    return lang, mc_lang, jars, tags, tags_of, loot, names, blocks, recipes, es_name


def acquisition(catalog_ids, es_name, loot, tags):
    """Dónde se consigue cada objeto en Diosesmon según los datos de la wiki y del mod."""
    sources = {item_id: [] for item_id in catalog_ids}
    by_name = {}
    for item_id in catalog_ids:
        by_name.setdefault(fold(es_name(item_id)), item_id)

    def match_name(text):
        key = fold(re.sub(r"(^\d[\d.]*\s*[x×]?\s*)|(\s*[x×]\s*\d+$)", "", str(text)))
        if key in NAME_ALIASES:
            return NAME_ALIASES[key]
        if key in by_name:
            return by_name[key]
        singular = " ".join(re.sub(r"(es|s)$", "", word) if len(word) > 3 else word for word in key.split())
        for name, item_id in by_name.items():
            if " ".join(re.sub(r"(es|s)$", "", word) if len(word) > 3 else word for word in name.split()) == singular:
                return item_id
        return None

    def add(item_id, entry):
        if item_id in sources and entry not in sources[item_id]:
            sources[item_id].append(entry)

    unmatched = []
    # Kits
    for kit in js_value("data/kits-data.js", "DIOSESMON_KITS") or []:
        if kit.get("available") is False:
            continue  # Kits fuera de rotación: ahora no se pueden comprar, no cuentan como «dónde se consigue».
        for item in kit["items"]:
            add(item["id"], {"t": "kit", "l": f"Kit {kit['name']}", "h": f"kits.html#{kit['key']}"})
    # Tesoros
    for box in js_value("data/tesoros-data.js", "DIOSESMON_TREASURES") or []:
        if box.get("kind") != "items":
            continue
        for reward in box.get("rewards", []):
            icon = reward.get("icon", "")
            match = re.match(r"assets/img/items/([^/]+)/([^/]+)\.png$", icon)
            item_id = f"{match.group(1)}:{match.group(2)}" if match else None
            if item_id not in sources:
                item_id = match_name(reward["name"])
            if item_id:
                add(item_id, {"t": "tesoro", "l": f"Tesoro {box['name']}", "h": f"tesoros.html#{box['key']}"})
            elif not re.search(r"dioses ?coins|pokedolares|\bdt\b|3x3|eevee|greninja", fold(reward["name"])):
                unmatched.append(f"Tesoro {box['name']}: {reward['name']}")
    # Drops de Pokémon (solo los del servidor y sin los desactivados)
    species = {entry["id"]: entry for entry in js_value("data/pokemon-data.js", "POKEDEX_SPECIES") if entry.get("kind") == "pokemon"}
    details = js_value("data/pokemon-details.js", "POKEDEX_DETAILS")
    overrides = js_object_loose("data/drop-overrides.js", "DIOSESMON_DROP_OVERRIDES")
    drops = {}
    for dex, detail in details.items():
        entry = species.get(int(dex))
        if not entry or not is_available(int(dex)):
            continue
        disabled = set((overrides.get(entry["identifier"]) or {}).get("disabled", []))
        for drop in (detail.get("drops") or {}).get("entries", []):
            if drop["item"] in sources and drop["item"] not in disabled:
                drops.setdefault(drop["item"], []).append([int(dex), entry["name"]])
    for item_id, mons in drops.items():
        add(item_id, {"t": "drop", "p": sorted(mons)})
    # Cazas (Poké Balls por PokéAPI)
    cazas = js_object_loose("data/cazas-data.js", "DIOSESMON_CAZAS")
    for kind in cazas.get("kinds", []):
        for tier in kind.get("tiers", []):
            for _, ball, _ in tier.get("balls", []):
                add("cobblemon:" + ball.replace("-", "_"), {"t": "caza", "l": kind["title"], "h": "guia.html#guide-cazas-title"})
    # Gimnasios y DexRewards (por nombre)
    gyms = js_object_loose("data/gimnasios-data.js", "GYM_REGIONS")
    for region in gyms:
        for leader in region.get("battles", []):
            for reward in leader.get("rewards", []):
                item_id = match_name(reward)
                if item_id:
                    add(item_id, {"t": "premio", "l": f"Gimnasio: {leader['name']}", "h": "gimnasios.html"})
    for section in js_object_loose("data/dexrewards-data.js", "DEXREWARDS"):
        for milestone in section.get("milestones", []):
            if not isinstance(milestone, dict):
                continue
            for reward in milestone.get("rewards", []):
                item_id = match_name(reward)
                if item_id:
                    add(item_id, {"t": "premio", "l": f"DexRewards: {section['title']} {milestone['percent']} %", "h": "dexrewards.html"})
    # Zona Safari
    add("cobblemon:safari_ball", {"t": "safari", "l": "Tienda de la Zona Safari (3 SP) y al entrar", "h": "safari.html"})
    for item_id, *_ in SAFARI_EXTRA:
        add(item_id, {"t": "safari", "l": "Tienda de la Zona Safari", "h": "safari.html"})
    # Mundo: cofres de estructuras, pesca, yacimientos y bloques que se pican.
    world = {}
    for key in loot:
        ns, _, path = key.partition(":")
        if path.startswith("blocks/"):
            block = ns + ":" + path[len("blocks/"):]
            if re.search(r"ore|crystal|cluster|meteorite|apricorn$|mushroom|gracidea", block) and "potted" not in block:
                for item_id in loot_items(loot, key):
                    if item_id != block:
                        world.setdefault(item_id, []).append(("bloque", f"Al picar o romper: {es_name(block)}", None))
            continue
        if path.startswith("fossils/"):
            label, href = "Yacimientos de fósiles (excavando)", "fosiles.html"
        elif path.startswith("fishing/") or path.startswith("injection/gameplay"):
            label, href = "Pesca con Poké Caña (tesoros)", "pesca.html"
        elif ns == "cobblemon" and (path.startswith("ruins/") or path.startswith("injection/") or path.startswith("villages/") or path.startswith("shipwreck_coves/")):
            label, href = "Cofres de estructuras (ruinas, aldeas, naufragios…)", None
        elif ns == "mega_showdown" and not path.startswith("sets/"):
            label, href = "Estructuras de Mega Showdown (observatorio, yacimientos…)", None
        else:
            continue
        for item_id in loot_items(loot, key):
            world.setdefault(item_id, []).append(("mundo", label, href))
    # Teralitos: Mega Showdown los da al vencer Pokémon (config del servidor: 10 % y 1 % el Astral).
    for item_id in catalog_ids:
        if item_id.startswith("mega_showdown:") and item_id.endswith("_tera_shard"):
            chance = "1 %" if "stellar" in item_id else "10 %"
            world.setdefault(item_id, []).append(("mundo", f"Al vencer Pokémon salvajes ({chance}, Mega Showdown)", None))
    for item_id, entries in world.items():
        seen = []
        for kind, label, href in entries:
            if label in seen:
                continue
            seen.append(label)
            entry = {"t": kind, "l": label}
            if href:
                entry["h"] = href
            add(item_id, entry)
    return sources, unmatched


# ---------------------------------------------------------------- Recetas del servidor

SERVER_KINDS = {"minecraft:crafting_shaped": "shaped", "minecraft:crafting_shapeless": "shapeless", "cobblemon:cooking_pot": "pot",
                "cobblemon:cooking_pot_shapeless": "potless", "minecraft:smelting": "smelt", "minecraft:blasting": "blast",
                "minecraft:smoking": "smoke", "minecraft:campfire_cooking": "campfire", "cobblemon:brewing_stand": "brew",
                "minecraft:stonecutting": "stone", "minecraft:smithing_transform": "smith"}


def import_server_dump(path):
    """Guarda en tools/source una copia reducida del volcado de /wikirecetas (solo lo de nuestros mods)."""
    dump = json.loads(Path(path).read_text(encoding="utf-8"))
    keep = set(ORIGINS)
    rows = []
    for row in dump.get("recipes", []):
        result = row.get("result")
        if not result or result.split(":")[0] not in keep or row.get("serializer") not in SERVER_KINDS:
            continue
        rows.append({key: row[key] for key in ("id", "serializer", "result", "count", "recipe", "width", "height", "ingredients", "error") if key in row})
    compact = {
        "recipeCount": dump.get("recipeCount"),  # Sin fecha ni dirección del servidor: el repo es público.
        "recipes": sorted(rows, key=lambda row: (row["result"], row["id"])),
        "notCrafted": sorted(item for item in dump.get("notCrafted", []) if item.split(":")[0] in keep),
    }
    SERVER_RECIPES.parent.mkdir(parents=True, exist_ok=True)
    SERVER_RECIPES.write_text(json.dumps(compact, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    detailed = sum(1 for row in rows if "recipe" in row or "ingredients" in row)
    print(f"Recetas del servidor importadas: {len(rows)} de nuestros mods ({detailed} con detalle) -> {SERVER_RECIPES.relative_to(ROOT)}")


def load_server_recipes(tags, es_name):
    """Recetas del servidor en el formato de la wiki. Devuelve (recetas por objeto, objetos sin receta, objetos con receta, fecha)."""
    if not SERVER_RECIPES.exists():
        return None
    dump = json.loads(SERVER_RECIPES.read_text(encoding="utf-8"))
    by_members = {}
    for tag in sorted(tags, key=lambda t: (t not in TAG_LABELS, len(t))):
        members = frozenset(resolve_tag(tags, tag))
        if len(members) > 1:
            by_members.setdefault(members, tag)

    def options_of(ingredient):
        if ingredient is None:
            return []
        if isinstance(ingredient, list):
            return [value for part in ingredient for value in options_of(part)]
        if isinstance(ingredient, dict):
            if "tag" in ingredient:
                return ["#" + ingredient["tag"]]
            value = ingredient.get("item") or ingredient.get("id")
            return [value] if value else []
        return [str(ingredient)] if ingredient else []

    def ref(ingredient):
        options = list(dict.fromkeys(options_of(ingredient)))
        if not options:
            return None
        if len(options) == 1:
            return options[0]
        tag = by_members.get(frozenset(options))
        if tag:
            return "#" + tag
        if len(options) <= 4:
            return "|".join(options)
        # Muchas opciones (por ejemplo, cualquier caña): se agrupan en una etiqueta propia de la wiki.
        synthetic = f"wiki:any_{options[0].split(':')[-1]}_{len(options)}"
        tags[synthetic] = options
        TAG_LABELS.setdefault(synthetic, (f"{es_name(options[0])} (o cualquiera de {len(options)} opciones)", options[0]))
        return "#" + synthetic

    recipes, produced = {}, set()
    for row in dump["recipes"]:
        result_id, kind = row["result"], SERVER_KINDS[row["serializer"]]
        produced.add(result_id)
        data = row.get("recipe")
        recipe = {"k": kind, "m": "diosesmon"}
        if row.get("count", 1) > 1:
            recipe["n"] = row["count"]
        if data is None:
            cells = [ref(options) for options in row.get("ingredients", [])]
            if not cells or all(cell is None for cell in cells):
                recipes.setdefault(result_id, [])  # Existe en el servidor pero sin detalle.
                continue
            if kind in ("shaped", "pot"):
                width = row.get("width") or (3 if len(cells) == 9 else len(cells))
                recipe["g"] = [cells[i:i + width] for i in range(0, len(cells), width)]
            else:
                counted = {}
                for cell in cells:
                    if cell:
                        counted[cell] = counted.get(cell, 0) + 1
                recipe["i"] = [[key, amount] for key, amount in counted.items()]
        elif kind in ("shapeless", "potless"):
            counted = {}
            for ingredient in data.get("ingredients", []):
                key = ref(ingredient)
                counted[key] = counted.get(key, 0) + 1
            recipe["i"] = [[key, amount] for key, amount in counted.items()]
            if data.get("seasoningTag") not in ("cobblemon:empty", None):
                recipe["s"] = True
        elif kind == "brew":
            recipe["i"] = [[ref(data.get("bottle")), 1], [ref(data.get("input")), 1]]
        elif kind == "smith":
            recipe["i"] = [[ref(data.get(part)), 1] for part in ("template", "base", "addition") if data.get(part)]
        else:
            recipe["i"] = [[ref(data.get("ingredient")), 1]]
        recipes.setdefault(result_id, [])
        if recipe not in recipes[result_id]:
            recipes[result_id].append(recipe)
    return recipes, set(dump.get("notCrafted", [])), produced, (dump.get("createdAt") or "")[:10]


# ---------------------------------------------------------------- Excel

def read_excel():
    if not EXCEL.exists():
        return {}
    book = load_workbook(EXCEL, data_only=True)
    sheet = book["Objetos"]
    header = [str(cell.value or "").strip() for cell in sheet[1]]
    col = {name: header.index(name) for name in HEADERS if name in header}
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item_id = str(row[col["ID"]] or "").strip()
        if not item_id:
            continue
        get = lambda name: str(row[col[name]] or "").strip() if name in col and row[col[name]] is not None else ""
        rows[item_id] = {"state": get("Crafteo en Diosesmon"), "where": get("Se consigue en"), "note": get("Nota")}
    return rows


def create_excel(items):
    existing = read_excel()
    book = Workbook()
    sheet = book.active
    sheet.title = "Objetos"
    sheet.append(HEADERS)
    labels = {key: label for key, label, _ in CATEGORIES}
    added = 0
    for item in items:
        old = existing.get(item["id"])
        if old is None:
            added += 1
            state = "Desactivado" if item["preset_off"] else ("Activo" if item["recipe"] else "Sin receta")
            where, note = "", item["preset_note"]
        else:
            state, where, note = old["state"] or ("Activo" if item["recipe"] else "Sin receta"), old["where"], old["note"]
        if item.get("server"):
            state = STATE_LABEL[item["server"]]  # Lo que dice el servidor (recetas-servidor.json) manda.
        sheet.append([item["id"], item["name"], labels[item["cat"]], ORIGINS[item["id"].split(":")[0]], "Sí" if item["recipe"] else "No", state, where, note])
    # Filas que ya no existen en los mods: se conservan al final para no perder lo escrito.
    known = {item["id"] for item in items}
    for item_id, old in existing.items():
        if item_id not in known:
            sheet.append([item_id, "(ya no está en los mods)", "", "", "No", old["state"], old["where"], old["note"]])

    help_sheet = book.create_sheet("Cómo usar")
    for line in [
        "Objetos de Diosesmon para la página Objetos de la wiki",
        "",
        "1. Hoja «Objetos»: una fila por objeto de Cobblemon, Mega Showdown, Navas ZA Megas y la tienda del Safari.",
        "2. «Crafteo en Diosesmon»: Activo (se puede craftear), Desactivado (el servidor lo tiene quitado) o Sin receta.",
        "   Si existe tools/source/recetas-servidor.json (sacado en el juego con /wikirecetas), el estado lo decide el servidor.",
        "   «Tiene receta en el mod» es informativo: dice si el mod trae receta, no lo cambies.",
        "3. «Se consigue en» (opcional): texto libre que sale en la wiki, por ejemplo «NPC Tienda del Centro Pokémon».",
        "   Los kits, tesoros, drops, cazas, gimnasios y el Safari ya salen solos: escribe aquí solo lo que falte.",
        "4. «Nota» (opcional): aviso que sale en la ficha del objeto (por ejemplo, por qué está desactivado).",
        "5. Guarda el Excel y ejecuta:  python -I tools/build_objetos_data.py",
        "   Eso genera data/objetos-data.js. Después haz commit y push para publicarlo.",
        "",
        "Si se actualizan los mods: python -I tools/build_objetos_data.py --crear (añade los objetos nuevos y conserva lo escrito).",
        "Truco: filtra la columna «Categoría» para ir por partes (Poké Balls, Medicinas…).",
    ]:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.column_dimensions["A"].width = 120

    header_fill = PatternFill("solid", fgColor="1F3A5C")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for index, width in enumerate([32, 30, 34, 16, 14, 22, 44, 50]):
        sheet.column_dimensions[chr(65 + index)].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    validation = DataValidation(type="list", formula1='"Activo,Desactivado,Sin receta"', allow_blank=False)
    sheet.add_data_validation(validation)
    cells = f"F2:F{sheet.max_row}"
    validation.add(cells)
    sheet.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=['"Desactivado"'], fill=PatternFill("solid", fgColor="F8C9D0")))
    sheet.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=['"Activo"'], fill=PatternFill("solid", fgColor="CDEFD9")))
    sheet.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=['"Sin receta"'], fill=PatternFill("solid", fgColor="E3E7EE")))
    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    book.save(EXCEL)
    print(f"Excel listo: {EXCEL.relative_to(ROOT)} ({len(items)} objetos, {added} filas nuevas)")


# ---------------------------------------------------------------- Principal

def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if "--recetas" in sys.argv:
        import_server_dump(sys.argv[sys.argv.index("--recetas") + 1])
    lang, mc_lang, jars, tags, tags_of, loot, names, blocks, recipes, es_name = build_catalog()

    # El servidor manda: si un objeto tiene receta allí se usa esa; si no la tiene, no se craftea.
    server = load_server_recipes(tags, es_name)
    server_state, server_items, server_date = {}, set(), ""
    if server:
        server_recipes, not_crafted, produced, server_date = server
        for item_id in names:
            if item_id in produced:
                server_state[item_id] = "on"
                if server_recipes.get(item_id):
                    recipes[item_id] = server_recipes[item_id]
                    server_items.add(item_id)
            elif item_id in not_crafted:
                server_state[item_id] = "off" if recipes.get(item_id) else "none"

    comp = js_value("data/competitivo-data.js", "DIOSESMON_COMP") or {}
    comp_effects = {row[0]: row[4] for row in comp.get("items", [])}
    stone_info = {}
    megas = js_value("data/megas-data.js", "DIOSESMON_MEGAS") or {"megas": []}
    for mega in megas["megas"]:
        if mega.get("stone"):
            stone_id = f"mega_showdown:{mega['stone']['id']}"
            for ns in ("mega_showdown", "zamega"):
                candidate = f"{ns}:{mega['stone']['id']}"
                if candidate in names:
                    stone_id = candidate
            text = f"Megapiedra: si la lleva {mega['base']}, puede convertirse en {mega['name']} en combate."
            stone_info[stone_id] = stone_info[stone_id].replace(" en combate.", f" o {mega['name']} en combate.") if stone_id in stone_info else text

    catalog_ids = sorted(item_id for item_id in names if item_id not in HIDDEN)
    extra_ids = [row[0] for row in SAFARI_EXTRA]
    sources, unmatched = acquisition(set(catalog_ids) | set(extra_ids), es_name, loot, tags)

    # Iconos (de los .jar; los de Minecraft para ingredientes se descargan una vez).
    indexes, icons = {}, {}
    used_refs = set()
    for item_id in catalog_ids:
        for recipe in recipes.get(item_id, []):
            cells = [cell for row in recipe.get("g", []) for cell in row] + [pair[0] for pair in recipe.get("i", [])]
            for cell in cells:
                for option in (cell or "").split("|"):
                    if option:
                        used_refs.add(option)
    for tag in [r[1:] for r in used_refs if r.startswith("#")]:
        icon_id = TAG_LABELS.get(tag, (None, None))[1] or next(iter(resolve_tag(tags, tag)), None)
        if icon_id:
            used_refs.add(icon_id)
    for item_id in set(catalog_ids) | {r for r in used_refs if not r.startswith("#")}:
        icons[item_id] = ensure_icon(item_id, jars, indexes)

    items = []
    for item_id in catalog_ids:
        cat = category_of(item_id, tags_of, item_id in blocks)
        if cat == "otros" and any(recipe["k"] in ("pot", "potless") for recipe in recipes.get(item_id, [])):
            cat = "cocina"  # Platos de la olla de campamento.
        if cat == "bayas":
            continue
        has_recipe = bool(recipes.get(item_id))
        # Bloques sin receta, sin icono y sin forma de conseguirse: son variantes internas (cultivos, cofres de Gimmighoul…).
        if not has_recipe and not icons.get(item_id) and not sources.get(item_id):
            continue
        iv_candy = "cobblemon:iv_candies" in tags_of.get(item_id, set())
        items.append({
            "id": item_id, "name": es_name(item_id), "cat": cat, "recipe": has_recipe,
            "preset_off": item_id in KNOWN_DISABLED or iv_candy,
            "preset_note": KNOWN_DISABLED.get(item_id) or (IV_CANDY_NOTE if iv_candy else ""),
            "server": server_state.get(item_id),
        })
    for item_id, name, cat, _, _ in SAFARI_EXTRA:
        items.append({"id": item_id, "name": name, "cat": cat, "recipe": False, "preset_off": False, "preset_note": "", "server": None})
    items.sort(key=lambda item: ([key for key, _, _ in CATEGORIES].index(item["cat"]), fold(item["name"])))

    if "--crear" in sys.argv or not EXCEL.exists():
        create_excel(items)
    sheet = read_excel()

    extra_desc = {row[0]: row[3] for row in SAFARI_EXTRA}
    extra_icon = {row[0]: row[4] for row in SAFARI_EXTRA}
    out_items, counts = [], {"on": 0, "off": 0, "none": 0}
    for item in items:
        item_id = item["id"]
        row = sheet.get(item_id, {})
        state = server_state.get(item_id) or STATES.get(fold(row.get("state", "")).replace("  ", " "))
        if not state:
            state = "off" if item["preset_off"] else ("on" if item["recipe"] else "none")
        counts[state] += 1
        entry = {"id": item_id, "n": item["name"], "c": item["cat"], "s": state}
        if item_id.split(":")[0] != "cobblemon":
            entry["o"] = item_id.split(":")[0]
        desc = extra_desc.get(item_id) or description(lang, item_id, comp_effects, stone_info)
        if desc:
            entry["d"] = desc
        icon = extra_icon.get(item_id) or icons.get(item_id) or ICON_FALLBACK.get(item_id)
        if icon:
            entry["i"] = icon
        if recipes.get(item_id):
            entry["r"] = recipes[item_id][:4]
            if item_id in server_items:
                entry["rs"] = 1  # Receta real del servidor.
        if sources.get(item_id):
            entry["g"] = sources[item_id]
        if row.get("where"):
            entry["x"] = row["where"]
        note = row.get("note") if item_id in sheet else item["preset_note"]
        if note:
            entry["note"] = note
        out_items.append(entry)

    # Nombres e iconos de los ingredientes que no están en la lista.
    listed = {entry["id"] for entry in out_items}
    ref_names = {}
    for ref in sorted(used_refs):
        if ref.startswith("#"):
            tag = ref[1:]
            label, icon_id = TAG_LABELS.get(tag, (None, None))
            members = resolve_tag(tags, tag)
            icon_id = icon_id or (members[0] if members else None)
            dye = re.fullmatch(r"c:dyes/(\w+)", tag)
            if dye:
                label, icon_id = f"Tinte {DYE_COLORS.get(dye.group(1), dye.group(1))}", f"minecraft:{dye.group(1)}_dye"
            ns, _, tag_path = tag.partition(":")
            label = label or clean(lang.get(f"tag.item.{ns}.{tag_path.replace('/', '.')}") or "")
            if not label:
                label = "Cualquier " + tag.split(":")[-1].split("/")[-1].replace("_", " ")
            ref_names[ref] = [label, icons.get(icon_id) or (ensure_icon(icon_id, jars, indexes) if icon_id else None)]
        elif ref not in listed:
            ref_names[ref] = [es_name(ref), icons.get(ref)]

    data = {
        "categories": [[key, label, emoji] for key, label, emoji in CATEGORIES],
        "origins": ORIGINS,
        "items": out_items,
        "refs": ref_names,
        "counts": {"total": len(out_items), **counts, "recipes": sum(1 for entry in out_items if entry.get("r"))},
    }
    dump = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    OUTPUT.write_text(
        "// Generado por tools/build_objetos_data.py desde Cobblemon 1.7.3, Mega Showdown, Navas ZA Megas, los datos de la wiki\n"
        "// y tools/source/objetos-diosesmon.xlsx (el estado del crafteo lo marca el staff ahí). No editar a mano.\n"
        "//   s: on (se craftea) · off (desactivado en Diosesmon) · none (sin receta) · r: recetas · g: dónde se consigue\n"
        f"window.DIOSESMON_OBJETOS = {dump};\n",
        encoding="utf-8",
    )
    page = ROOT / "objetos.html"
    if page.exists():
        html = page.read_text(encoding="utf-8")
        page.write_text(re.sub(r"objetos-data\.js\?v=(\d+)", lambda m: f"objetos-data.js?v={int(m.group(1)) + 1}", html), encoding="utf-8")
    by_cat = {}
    for entry in out_items:
        by_cat[entry["c"]] = by_cat.get(entry["c"], 0) + 1
    print(f"data/objetos-data.js: {len(out_items)} objetos ({counts['on']} activos, {counts['off']} desactivados, {counts['none']} sin receta)")
    if server:
        on_without = [e["id"] for e in out_items if e["s"] == "on" and not e.get("rs")]
        print(f"  recetas del servidor ({server_date}): {len(server_items)} objetos; activos sin detalle de receta: {len(on_without)}", on_without[:30])
    print("  por categoría:", by_cat)
    print("  sin icono:", sum(1 for e in out_items if not e.get("i")), [e["id"] for e in out_items if not e.get("i")][:40])
    print("  sin descripción:", sum(1 for e in out_items if not e.get("d")))
    print("  ingredientes sin icono:", [k for k, v in ref_names.items() if not v[1]][:40])
    if unmatched:
        print("  premios sin objeto reconocido:", unmatched)


if __name__ == "__main__":
    main()
