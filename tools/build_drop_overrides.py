"""Drops de Diosesmon editables en Excel.

  python tools/build_drop_overrides.py --crear   Crea (o completa) tools/source/drops-diosesmon.xlsx con todos los
                                                 drops de Cobblemon 1.7.3, conservando lo que ya hayas marcado.
  python tools/build_drop_overrides.py           Lee el Excel y genera data/drop-overrides.js para la Pokédex.

En el Excel:
- Hoja «Drops»: una fila por Pokémon y objeto. Pon «No» en «Activo en Diosesmon» si el servidor lo tiene desactivado.
- Hoja «Por objeto»: pon «Sí» en «Desactivar para todos» para quitar un objeto a todos los Pokémon de golpe.
- «Nota» (opcional): texto que sale bajo la tabla de drops de ese Pokémon en la wiki.
"""
import json
import re
import sys
import tarfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
EXCEL = ROOT / "tools" / "source" / "drops-diosesmon.xlsx"
OUTPUT = ROOT / "data" / "drop-overrides.js"
CACHE = ROOT / "tools" / "build-cache" / "cobblemon-1.7.3-data.tar.gz"
MAX_NATIONAL_ID = 493
DROP_HEADERS = ["Nº", "Pokémon", "Identificador", "Objeto", "ID del objeto", "Tipo", "Cantidad / prob.", "Activo en Diosesmon", "Nota (opcional)"]
ITEM_HEADERS = ["Objeto", "ID del objeto", "Tipo", "Lo sueltan", "Desactivar para todos"]


def js_value(path, name):
    source = (ROOT / path).read_text(encoding="utf-8")
    return json.loads(re.search(rf"window\.{name} = (.*?);\n", source, re.S).group(1))


def held_items():
    """Objetos equipables de Cobblemon (etiquetas item/held)."""
    held = set()
    if not CACHE.exists():
        return held
    with tarfile.open(CACHE) as tar:
        for member in tar.getmembers():
            name = member.name.split("/data/")[-1]
            if member.isfile() and name.startswith("cobblemon/tags/item/held"):
                for value in json.loads(tar.extractfile(member).read()).get("values", []):
                    held.add(value if isinstance(value, str) else value.get("id"))
    return held


def item_type(item_id, held):
    if item_id.startswith("minecraft:"):
        return "Minecraft"
    if item_id.endswith("_berry"):
        return "Baya"
    if item_id in held:
        return "Objeto equipable"
    return "Cobblemon (otro)"


def amount(entry):
    parts = []
    if "quantity" in entry:
        parts.append(f"x{entry['quantity']}")
    if "chance" in entry:
        parts.append(f"{entry['chance']:g}%")
    return " · ".join(parts) or "x1"


def current_overrides():
    """Lo que ya está en drop-overrides.js (para no perderlo al crear el Excel)."""
    if not OUTPUT.exists():
        return {}
    text = OUTPUT.read_text(encoding="utf-8")
    body = text[text.index("{"):text.rindex("}") + 1]
    body = re.sub(r"//.*", "", body)
    body = re.sub(r"([{,]\s*)([A-Za-z_][\w-]*)\s*:", r'\1"\2":', body)
    body = re.sub(r",(\s*[}\]])", r"\1", body)
    return json.loads(body)


def create():
    species = json.loads(re.search(r"window\.POKEDEX_SPECIES = (\[.*?\]);", (ROOT / "data" / "pokemon-data.js").read_text(encoding="utf-8"), re.S).group(1))
    details = js_value("data/pokemon-details.js", "POKEDEX_DETAILS")
    held = held_items()

    # Lo ya marcado: primero el Excel existente, si no, drop-overrides.js.
    marked, notes, global_off = {}, {}, set()
    if EXCEL.exists():
        book = load_workbook(EXCEL)
        for row in book["Drops"].iter_rows(min_row=2, values_only=True):
            if row[2] and row[4]:
                marked[(row[2], row[4])] = row[7]
                if row[8]:
                    notes[row[2]] = row[8]
        for row in book["Por objeto"].iter_rows(min_row=2, values_only=True):
            if row[1] and str(row[4] or "").strip().lower() in ("sí", "si"):
                global_off.add(row[1])
    else:
        for identifier, override in current_overrides().items():
            for item in override.get("disabled", []):
                marked[(identifier, item)] = "No"
            if override.get("note"):
                notes[identifier] = override["note"]

    book = Workbook()
    sheet = book.active
    sheet.title = "Drops"
    sheet.append(DROP_HEADERS)
    items = {}
    for entry in species:
        if entry.get("kind") != "pokemon" or entry["id"] > MAX_NATIONAL_ID or not entry.get("available", True):
            continue
        drops = (details.get(str(entry["id"])) or {}).get("drops") or {}
        for index, drop in enumerate(drops.get("entries", [])):
            identifier = entry["identifier"]
            kind = item_type(drop["item"], held)
            active = marked.get((identifier, drop["item"]), "Sí")
            sheet.append([entry["id"], entry["name"], identifier, drop["name"], drop["item"], kind, amount(drop), active,
                          notes.get(identifier, "") if index == 0 else ""])
            info = items.setdefault(drop["item"], {"name": drop["name"], "type": kind, "count": 0})
            info["count"] += 1

    item_sheet = book.create_sheet("Por objeto")
    item_sheet.append(ITEM_HEADERS)
    order = {"Objeto equipable": 0, "Cobblemon (otro)": 1, "Baya": 2, "Minecraft": 3}
    for item_id, info in sorted(items.items(), key=lambda pair: (order[pair[1]["type"]], pair[1]["name"])):
        item_sheet.append([info["name"], item_id, info["type"], info["count"], "Sí" if item_id in global_off else "No"])

    help_sheet = book.create_sheet("Cómo usar")
    for line in [
        "Drops de Diosesmon para la Pokédex de la wiki",
        "",
        "1. Hoja «Drops»: pon «No» en «Activo en Diosesmon» para cada drop que el servidor tiene desactivado.",
        "2. Hoja «Por objeto»: pon «Sí» en «Desactivar para todos» para quitar ese objeto a todos los Pokémon a la vez",
        "   (por ejemplo Semilla Milagro, Carbón o Agua Mística). Filtra la columna «Tipo» por «Objeto equipable».",
        "3. «Nota (opcional)»: texto que aparece bajo la tabla de drops de ese Pokémon en la wiki (basta en una fila).",
        "4. Guarda el Excel y ejecuta:  python tools/build_drop_overrides.py",
        "   Eso genera data/drop-overrides.js. Después haz commit y push para publicarlo.",
        "",
        "Si actualizas los datos de Cobblemon: python tools/build_drop_overrides.py --crear (conserva lo que marcaste).",
    ]:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.column_dimensions["A"].width = 120

    style_sheets(sheet, item_sheet)
    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    book.save(EXCEL)
    print(f"Excel listo: {EXCEL.relative_to(ROOT)} ({sheet.max_row - 1} drops, {len(items)} objetos)")


def style_sheets(sheet, item_sheet):
    header_fill = PatternFill("solid", fgColor="1F3A5C")
    red = PatternFill("solid", fgColor="F8C9D0")
    green = PatternFill("solid", fgColor="CDEFD9")
    for target, widths, choice_col in ((sheet, [6, 16, 16, 24, 30, 18, 16, 20, 40], "H"), (item_sheet, [26, 32, 20, 12, 22], "E")):
        for cell in target[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for index, width in enumerate(widths):
            target.column_dimensions[chr(65 + index)].width = width
        target.freeze_panes = "A2"
        target.auto_filter.ref = target.dimensions
        validation = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False)
        target.add_data_validation(validation)
        cells = f"{choice_col}2:{choice_col}{target.max_row}"
        validation.add(cells)
        bad, good = ('"No"', '"Sí"') if target is sheet else ('"Sí"', '"No"')
        target.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=[bad], fill=red))
        target.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=[good], fill=green))


def build():
    if not EXCEL.exists():
        sys.exit("No existe el Excel: créalo con  python tools/build_drop_overrides.py --crear")
    book = load_workbook(EXCEL, data_only=True)
    global_off = {row[1] for row in book["Por objeto"].iter_rows(min_row=2, values_only=True)
                  if row[1] and str(row[4] or "").strip().lower() in ("sí", "si")}
    overrides = {}
    for row in book["Drops"].iter_rows(min_row=2, values_only=True):
        identifier, item_id = row[2], row[4]
        if not identifier or not item_id:
            continue
        entry = overrides.setdefault(identifier, {"disabled": [], "note": ""})
        if str(row[7] or "").strip().lower() == "no" or item_id in global_off:
            if item_id not in entry["disabled"]:
                entry["disabled"].append(item_id)
        if row[8]:
            entry["note"] = str(row[8]).strip()
    lines = [
        "// Generado por tools/build_drop_overrides.py desde tools/source/drops-diosesmon.xlsx. Edita el Excel, no este archivo.",
        "//",
        "// Clave: identificador inglés del Pokémon. disabled: objetos desactivados en Diosesmon (salen tachados).",
        "// note: texto opcional bajo la tabla de drops.",
        "window.DIOSESMON_DROP_OVERRIDES = {",
    ]
    total = 0
    for identifier, entry in overrides.items():
        if not entry["disabled"] and not entry["note"]:
            continue
        parts = []
        if entry["disabled"]:
            parts.append(f"disabled: {json.dumps(entry['disabled'])}")
            total += len(entry["disabled"])
        if entry["note"]:
            parts.append(f"note: {json.dumps(entry['note'], ensure_ascii=False)}")
        key = identifier if re.fullmatch(r"[A-Za-z_]\w*", identifier) else json.dumps(identifier)
        lines.append(f"  {key}: {{ {', '.join(parts)} }},")
    lines.append("};")
    OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Sube la versión en pokedex.html para que los navegadores no usen la copia vieja.
    page = ROOT / "pokedex.html"
    html = page.read_text(encoding="utf-8")
    html = re.sub(r"drop-overrides\.js\?v=(\d+)", lambda match: f"drop-overrides.js?v={int(match.group(1)) + 1}", html)
    page.write_text(html, encoding="utf-8")
    print(f"data/drop-overrides.js generado: {total} drops desactivados en {sum(1 for e in overrides.values() if e['disabled'])} Pokémon")


if __name__ == "__main__":
    create() if "--crear" in sys.argv else build()
